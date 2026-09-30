#!/usr/bin/env python3
"""make_ESM_variance_components.py -- Online Resource (ESM number fixed later): variance components and
broad-sense heritability of the eight traits of the genotype x season model. Replaces Table 1 of the main text.

WHY (S. Hübner's comment on Table 1, 2026-09-30: "make this a full table with all variation components and put
it in the supplementaries"; user decision): Table 1 (H² only) leaves the main text; this table gives, for each of
the 8 traits, the variance and % of the total for genetic (V_G), season & block (V_E), genotype x season (V_GxE)
and residual (V_R), the total variance and H².

NOTHING IS RE-ESTIMATED. Values are copied from the step-00 outputs (read-only):
  00_THIN_Generate_Plots_For_Publication/outputs/subsection_1/tables/Variance_components_GxE.csv   variances, % (Fig. 1a)
  00_THIN_Generate_Plots_For_Publication/outputs/subsection_1/tables/Variance_components_GxE_WIDE.csv  same, wide (cross-check)
  00_THIN_Generate_Plots_For_Publication/outputs/subsection_1/tables/H2_GxE.csv                      H² (former Table 1)
written by 00_THIN_.../02_phenotypic_analysis_GxE_v2_streamlined.R, herit_calc_GxE():
lmer(trait ~ (1 | short_Tag) + (1 | season_Bl) + (1 | short_Tag:season)); H² = V_G / (V_G + V_E + V_GxE + V_R).

Checks (the script stops if any fails): every % recomputed from the variances equals the step-00 Percentage
(rounded to 2 decimals) and the WIDE table; the four % sum to 100; H² recomputed from the variances equals
H2_GxE.csv (3 decimals) and the genetic % / 100; the H² values equal those of the former Table 1 of the
manuscript (starch 0.474, fiber 0.248, protein 0.242, beta-glucan 0.234, flowering time 0.747, spike length 0.344,
grain weight 0.223, tillers < 0.001).
Tillers: V_G = 2.0e-08 of a total 76.26, so H² is written "< 0.001", as in the former Table 1 and the Results
(user decision 2026-09-28: an estimate is never exactly zero).

Outputs (this folder's ../tables/): ESM_variance_components.xlsx (sheet 1 the table, sheet 2 notes and sources),
ESM_variance_components.tsv (same values, for checking and diffs).
The ESM title block TAG asks for (article title, journal, authors, corresponding author) is added when the Online
Resources are numbered and assembled, not here.

Created 2026-09-30 (S. Hübner's structural comments; hand-over build/STRUCTURE_CHANGES_2026-10.md).
Run: python3 make_ESM_variance_components.py   (seconds; openpyxl)
"""
import csv, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path("/mnt/data/shahar/gwas_barley/morexV3_analysis/barley_for_publication_project")
TAB = ROOT / "00_THIN_Generate_Plots_For_Publication/outputs/subsection_1/tables"
OUT = pathlib.Path(__file__).resolve().parent.parent / "tables"
OUT.mkdir(exist_ok=True)

COMP = ["Genetics", "Season & Block", "G x E", "Residual"]
# row order of the former Table 1: nutritional, then morphological, each by decreasing H²
ORDER = ["Starch", "Fiber", "Protein", "β-glucan", "Flowering time", "Spike length", "Grain weight", "Tillers"]
GROUP = {t: "Nutritional" for t in ORDER[:4]} | {t: "Morphological" for t in ORDER[4:]}
TABLE1 = {"Starch": "0.474", "Fiber": "0.248", "Protein": "0.242", "β-glucan": "0.234", "Flowering time": "0.747",
          "Spike length": "0.344", "Grain weight": "0.223", "Tillers": "< 0.001"}   # manuscript Table 1, 2026-09-30


def rd(path):
    return list(csv.DictReader(open(path, encoding="utf-8")))


long = rd(TAB / "Variance_components_GxE.csv")
wide = {r["TraitLabel"]: r for r in rd(TAB / "Variance_components_GxE_WIDE.csv")}
h2 = {r["Trait"]: r for r in rd(TAB / "H2_GxE.csv")}
var = {}
for r in long:
    var.setdefault(r["TraitLabel"], {})[r["Component"]] = (float(r["Variance"]), float(r["Percentage"]),
                                                           float(r["TotalVar"]))
assert sorted(var) == sorted(ORDER) == sorted(wide) == sorted(h2), "trait sets differ"

rows = []
for t in ORDER:
    v = {c: var[t][c][0] for c in COMP}
    tot = sum(v.values())
    assert all(abs(var[t][c][2] - tot) < 1e-9 * max(1, tot) for c in COMP), f"{t}: TotalVar differs from the sum"
    pct = {c: 100 * v[c] / tot for c in COMP}
    for c, wcol in zip(COMP, ["Genetics_pct", "Season_Block_pct", "GxE_pct", "Residual_pct"]):
        assert round(pct[c], 2) == var[t][c][1] == float(wide[t][wcol]), f"{t} {c}: % differs"
    assert abs(sum(var[t][c][1] for c in COMP) - 100) < 0.02, f"{t}: % do not sum to 100"
    H2 = v["Genetics"] / tot
    assert round(H2, 3) == float(h2[t]["H2_GxE"]), f"{t}: H² differs from H2_GxE.csv"
    assert abs(round(H2 * 100, 2) - var[t]["Genetics"][1]) < 1e-9, f"{t}: H² != genetic % / 100"
    h2_cell = "< 0.001" if H2 < 0.0005 else f"{H2:.3f}"
    assert h2_cell == TABLE1[t], f"{t}: H² {h2_cell} differs from the former Table 1 ({TABLE1[t]})"
    rows.append([t, GROUP[t], v["Genetics"], pct["Genetics"], v["Season & Block"], pct["Season & Block"],
                 v["G x E"], pct["G x E"], v["Residual"], pct["Residual"], tot,
                 h2_cell if h2_cell.startswith("<") else round(H2, 3)])

HEAD = ["Trait", "Trait group", "V_G", "V_G (%)", "V_E (season & block)", "V_E (%)", "V_G×E", "V_G×E (%)",
        "V_R", "V_R (%)", "Total variance", "H²"]

# ---- TSV (full precision) ----------------------------------------------------
with open(OUT / "ESM_variance_components.tsv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f, delimiter="\t")
    w.writerow(HEAD)
    w.writerows(rows)

# ---- XLSX --------------------------------------------------------------------
CAPTION = ("Variance components and broad-sense heritability (H²) of four nutritional and four morphological "
           "traits, from a linear mixed model with genotype, season-by-block and genotype-by-season random effects "
           "fitted to the season-centered values of each trait (290 accessions). For each "
           "component, the variance and its percentage of the total phenotypic variance. V_G, genetic; V_E, "
           "environmental (season and block); V_G×E, genotype-by-season interaction; V_R, residual; "
           "H² = V_G / (V_G + V_E + V_G×E + V_R). Variances are in the squared units of each trait. Traits within each "
           "group are ordered by decreasing H².")
wb = Workbook()
ws = wb.active
ws.title = "Variance components"
bold = Font(bold=True)
thin = Side(style="thin")
ws.append(HEAD)
for c in ws[1]:
    c.font = bold
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(bottom=thin)
for r in rows:
    ws.append(r)
for row in ws.iter_rows(min_row=2):
    for c in row[2:11]:
        hdr = HEAD[c.column - 1]
        c.number_format = "0.00" if hdr.endswith("(%)") else "0.000E+00" if c.value < 0.001 else "0.0000"
    row[11].number_format = "0.000"
    row[11].alignment = Alignment(horizontal="right")
for i, wdt in enumerate([15, 14, 11, 9, 13, 9, 11, 10, 11, 9, 12, 8], 1):
    ws.column_dimensions[get_column_letter(i)].width = wdt
ws.row_dimensions[1].height = 30
ws.freeze_panes = "B2"

wn = wb.create_sheet("Notes")
notes = [("Caption (draft)", CAPTION),
         ("Model", "lme4: trait ~ (1 | genotype) + (1 | season:block) + (1 | genotype:season); variance components "
                   "from VarCorr; H² following Holland et al. (2003) as in the Materials and methods."),
         ("Tillers", "V_G = 2.0 × 10⁻⁸ of a total variance of 76.26, so H² is given as < 0.001."),
         ("Source", "00_THIN_Generate_Plots_For_Publication/outputs/subsection_1/tables/Variance_components_GxE.csv, "
                    "Variance_components_GxE_WIDE.csv, H2_GxE.csv (script 02_phenotypic_analysis_GxE_v2_streamlined.R, "
                    "herit_calc_GxE())"),
         ("Built by", "10_USED_Paper_writing/new_publishing_paper/supplementary/scripts/make_ESM_variance_components.py"),
         ("To add", "TAG ESM title block (article title, journal, authors, corresponding author's affiliation and "
                    "e-mail) when the Online Resources are numbered.")]
for k, v in notes:
    wn.append([k, v])
    wn.cell(wn.max_row, 1).font = bold
    wn.cell(wn.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    wn.cell(wn.max_row, 1).alignment = Alignment(vertical="top")
wn.column_dimensions["A"].width = 16
wn.column_dimensions["B"].width = 110
wb.save(OUT / "ESM_variance_components.xlsx")

print(f"wrote {OUT / 'ESM_variance_components.xlsx'} and .tsv ({len(rows)} traits; all checks passed)")
for r in rows:
    print(f"  {r[0]:15s} G {r[3]:6.2f}%  E {r[5]:6.2f}%  GxE {r[7]:6.2f}%  R {r[9]:6.2f}%  total {r[10]:9.4f}  H² {r[11]}")
