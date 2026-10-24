#!/usr/bin/env python3
"""package_for_review.py -- the supplement for S. Hübner's review (his request of 2026-10-04: all the material in one
Excel file, as sheets, and one Word file, each table and figure with its legend, as in E. Potapenko's papers; the PDFs
could not be edited). Writes ../for_review_Hubner/:

  Supplementary_Tables.xlsx                                 Tables S1-S22 = a copy of ../ESM_1.xlsx (build_ESM_tables.py)
  Supplementary_Figures.docx                                Figs. S1-S5, each with its legend (as ../ESM_2.pdf)
  Raw_phenotype_checks_uncited_in_manuscript_tables.xlsx    the unnumbered supplement (../unreferenced/), tables
  Raw_phenotype_checks_uncited_in_manuscript_figures.docx   the unnumbered supplement, figures
  Supplementary_for_review.zip                              the four files

RENUMBERED 2026-10-17 (S. Hübner's comment, user decision): the cited material is numbered as in the manuscript,
Supplementary Tables S1-S22 and Supplementary Figures S1-S5 (it was "Online Resources 1-13", files
Online_Resources_tables.xlsx and Online_Resources_figures.docx). The tables workbook is no longer assembled here: the
published ESM_1.xlsx already has this layout (one sheet per table, caption and column definitions above the table) and
is copied unchanged. The Word file is built from build_ESM_pdf.figs(), the same figures, panels and legends as ESM_2.pdf.
The xlsx helpers that lived here moved to esm_assemble.py (also used by build_ESM_tables.py).
The unnumbered supplement carries descriptive sheet and figure names instead of numbers (user, 2026-10-04: no
numbers, and no statement inside the files that they are not cited; "uncited_in_manuscript" is in the file names).

Figures: vector PDF pages are rendered at 300 dpi (pdftoppm); the 600-dpi PNG figures are downscaled to 300 dpi.
The unnumbered figures are captured by running build_ESM_extra.py's page functions with the LaTeX steps replaced.
The TAG submission package (../for_submission/) is not touched.

Requires python-docx (1.2.0, installed with pip --target in /mnt/data/shahar/gwas_barley/tools/pylib, 2026-10-04).
Run: python3 package_for_review.py   (the last step of build_all.sh since 2026-10-17). Created 2026-10-04.
"""
import math, os, re, shutil, subprocess, sys, tempfile, zipfile, pathlib
from copy import copy

sys.path.insert(0, "/mnt/data/shahar/gwas_barley/tools/pylib")
os.environ["TMPDIR"] = "/mnt/data/shahar/.tmp"

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import esm_common as C
from esm_assemble import md_parts, plain, rich, parse_about, copy_table_sheet, contents_sheet, check_sheet_refs
import build_ESM_pdf as BP
import build_ESM_extra as BX

OUT = C.SUPP / "for_review_Hubner"
POPPLER = BP.POPPLER
XLSX_CITED = "Supplementary_Tables.xlsx"         # = ESM_1.xlsx (Tables S1-S22), copied
DOCX_CITED = "Supplementary_Figures.docx"        # Figs. S1-S5, each with its legend (same figures and legends as ESM_2.pdf)
XLSX_EXTRA = "Raw_phenotype_checks_uncited_in_manuscript_tables.xlsx"
DOCX_EXTRA = "Raw_phenotype_checks_uncited_in_manuscript_figures.docx"
ZIP = "Supplementary_for_review.zip"

# the unnumbered supplement: descriptive names instead of Fig. E1-E7 and of the short sheet names
FIG_NAME = {"E1": "Nutritional traits by season", "E2": "Nutritional traits by block",
            "E3": "Correlations between seasons", "E4": "Correlations between traits",
            "E5": "Correlations between blocks", "E6": "Distributions of raw values",
            "E7": "Distributions of genotype BLUPs"}
EXTRA_SHEET = {"Summary by season": "Raw values by season", "Summary by block": "Raw values by block",
               "Correlations seasons": "Correlations between seasons", "Correlations traits": "Correlations between traits",
               "Correlations blocks": "Correlations between blocks", "BLUP summary": "Genotype BLUPs"}


def unnumber(s):
    """Fig. E1 / Figs. E1, E6 / Figs. E3–E5 -> the figure names of the unnumbered supplement"""
    def rep(m):
        ids = []
        for part in re.split(r", | and ", m.group(1)):
            if "–" in part:
                a, b = (int(x[1:]) for x in part.split("–"))
                ids += [f"E{k}" for k in range(a, b + 1)]
            else:
                ids.append(part)
        q = [f"“{FIG_NAME[i]}”" for i in ids]
        return f"the figure {q[0]}" if len(q) == 1 else "the figures " + ", ".join(q[:-1]) + " and " + q[-1]
    s = re.sub(r"Figs?\. (E\d(?:(?:, |–| and )E\d)*)", rep, s).replace("(the figure", "(figure")
    assert not re.search(r"\bE\d\b", s), s
    return s


def build_extra_xlsx():
    src, _label, _caption, sheets, defs, notes = parse_about(C.SUPP / "unreferenced/ESM_extra_phenotype_checks.xlsx")
    assert [s for s, _ in sheets] == list(EXTRA_SHEET), sheets
    wb = openpyxl.Workbook()
    panel = unnumber(BX.PANEL)
    for s, desc in sheets:
        legend = [([(f"{EXTRA_SHEET[s]}. ", True, False), (unnumber(desc), False, False)], False),
                  ([(panel, False, False)], False)]
        if defs.get(s):
            legend += [([("Columns", True, False)], True)]
            legend += [([(f"{col}: ", True, False), (unnumber(d), False, False)], False) for col, d in defs[s]]
        assert not notes
        copy_table_sheet(wb, src[s], EXTRA_SHEET[s], legend)
    contents_sheet(wb, "Raw phenotype checks: tables", [
        "Descriptive checks of the raw phenotypic data of the common-garden experiment (three seasons, 2019–20 to "
        "2021–22), one sheet per table. The figures are in the accompanying Word file.", panel],
        [("", [(unnumber(d), False, False)], [(EXTRA_SHEET[s], None)]) for s, d in sheets],
        col_b="Content", widths=(4, 95, 30))
    check_sheet_refs(wb)
    C.check_no_screen_categories(wb)
    wb.save(OUT / XLSX_EXTRA)
    print("wrote", OUT / XLSX_EXTRA, f"({len(wb.sheetnames) - 1} table sheets)")


# ---- figures: capture what the PDF builders put on their pages -----------------------------------------------------
def capture(module, call):
    rec = {"legend": None, "panels": []}
    saved = {k: getattr(module, k) for k in ("title_page", "figure_page", "compile_pages", "finish")}
    module.title_page = lambda n, legend, pages: rec.update(legend=legend, n=n) or ""
    module.figure_page = lambda n, label, graphic, ident, paper, opts: rec["panels"].append(
        dict(label=label, graphic=graphic, ident=ident, landscape="landscape" in paper)) or ""
    module.compile_pages = lambda *a, **k: []
    module.finish = lambda *a, **k: None
    try:
        call()
    finally:
        for k, v in saved.items():
            setattr(module, k, v)
    return rec


def render(graphic, work, tag):
    """figure -> PNG at about 300 dpi; returns (path, width px, height px)"""
    path, page = graphic
    path = pathlib.Path(path)
    out = work / f"{tag}.png"
    if path.suffix.lower() == ".pdf":
        subprocess.run([f"{POPPLER}/pdftoppm", "-png", "-r", "300", "-singlefile", "-f", str(page or 1),
                        "-l", str(page or 1), str(path), str(out.with_suffix(""))], check=True)
    else:
        im = Image.open(path)
        dpi = (im.info.get("dpi") or (300, 300))[0]
        if dpi > 310:
            im = im.resize((round(im.width * 300 / dpi), round(im.height * 300 / dpi)), Image.LANCZOS)
        im.save(out, dpi=(300, 300))
    w, h = Image.open(out).size
    return out, w, h


# ---- Word ----------------------------------------------------------------------------------------------------------
PORTRAIT = (21.0, 29.7)
MARGIN = 2.0


def new_doc(title):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(4)
    for name, size in (("Heading 1", 12),):
        h = doc.styles[name]
        h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
        rf = h.element.rPr.rFonts
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        rf.set(qn("w:ascii"), "Arial"); rf.set(qn("w:hAnsi"), "Arial")
        h.paragraph_format.space_before = Pt(0); h.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(PORTRAIT[0]), Cm(PORTRAIT[1])
    for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, a, Cm(MARGIN))
    page_number(sec)
    doc.core_properties.title = title
    doc.core_properties.author = C.AUTHORS
    doc.core_properties.comments = ""
    return doc


def page_number(sec):
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r)
    p._p.append(fld)


def add_par(doc, parts, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=None):
    p = doc.add_paragraph()
    p.alignment = align
    for t, b, i in parts:
        r = p.add_run(t); r.bold = b or None; r.italic = i or None
        if size:
            r.font.size = Pt(size)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def title_page(doc, heading, intro, items):
    p = add_par(doc, [(t, True, it) for t, it in C.TITLE_PARTS], size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
    for line in (C.JOURNAL, C.AUTHORS, C.CORR_LINE):
        add_par(doc, [(line, False, False)], align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    doc.add_paragraph(heading, style="Heading 1")
    for x in intro:
        add_par(doc, md_parts(x) if isinstance(x, str) else x)
    for parts in items:
        q = add_par(doc, parts, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3)
        q.paragraph_format.left_indent = Cm(0.5)


def est_lines(text, width_cm, size=10):
    return max(1, math.ceil(len(text) / (width_cm * 4.9 * 10 / size)))


def new_section(doc, want):
    """new page with the given orientation. python-docx puts the closing sectPr of the previous section in a new
    empty paragraph, which can spill onto a blank page: move it into the paragraph before"""
    doc.add_section(WD_SECTION.NEW_PAGE)
    body = doc.element.body
    holder = body.findall(qn("w:p"))[-1]
    prev = holder.getprevious()
    sp = holder.find(qn("w:pPr") + "/" + qn("w:sectPr"))
    if sp is not None and prev is not None and prev.tag == qn("w:p") and not holder.findall(".//" + qn("w:t")):
        prev.get_or_add_pPr().append(sp)
        body.remove(holder)
    sec = doc.sections[-1]
    sec.orientation = want
    w, h = (PORTRAIT[1], PORTRAIT[0]) if want == WD_ORIENT.LANDSCAPE else PORTRAIT
    sec.page_width, sec.page_height = Cm(w), Cm(h)


def add_figures(doc, heading, panels, caption, work, tag):
    """one page per panel (landscape where the PDF builder used landscape); heading on the first page; each panel's
    identification line under it; the caption (list of paragraphs, each a list of parts) after the last panel"""
    for k, pnl in enumerate(panels):
        want = WD_ORIENT.LANDSCAPE if pnl["landscape"] else WD_ORIENT.PORTRAIT
        brk = doc.sections[-1].orientation == want          # same orientation: page break; else a new section
        if not brk:
            new_section(doc, want)
        sec = doc.sections[-1]
        first = None
        tw = sec.page_width.cm - 2 * MARGIN
        th = sec.page_height.cm - 2 * MARGIN - 0.8
        if k == 0 and heading:
            first = doc.add_paragraph(heading, style="Heading 1")
            th -= 1.0
        ident = pnl["ident"] if (len(panels) > 1 or pnl["label"]) else ""
        if ident:
            th -= 0.45 * est_lines(plain(ident), tw, 9 if len(ident) <= 200 else 10) + 0.3
        if k == len(panels) - 1:
            need = sum(0.42 * est_lines(plain("".join(t for t, _, _ in c)), tw) + 0.2 for c in caption)
            if th - need >= 0.8 * th:                    # the legend fits under the figure without shrinking it
                th -= need                                  # by more than 20%; otherwise it flows to the next page
        png, wpx, hpx = render(pnl["graphic"], work, f"{tag}_{k}")
        w = min(tw, th * wpx / hpx)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(png), width=Cm(w))
        (first or p).paragraph_format.page_break_before = brk
        if ident and len(ident) > 200:
            add_par(doc, md_parts(ident))
        elif ident:
            add_par(doc, md_parts(ident), size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
    for c in caption:
        add_par(doc, c)


def build_cited_docx(work):
    """Figs. S1-S5 from build_ESM_pdf.figs(): the same figures, panels and legends as ESM_2.pdf"""
    spec = BP.figs()
    doc = new_doc(f"{C.TITLE_PLAIN} - Supplementary Figures")
    title_page(doc, "Supplementary Figures",
               [f"Supplementary Figures S1–S{C.N_FIGURES} of the manuscript, numbered in the order in which they are first "
                f"cited, each followed by its legend. The tables (Tables S1–S{C.N_TABLES}) are in the Excel file {XLSX_CITED}."],
               [[(f"Fig. S{k} ", True, False)] + md_parts(t) for k, (_, t) in enumerate(C.SUPP_FIGURES, 1)])
    for k, (key, _) in enumerate(C.SUPP_FIGURES, 1):
        f = spec[key]
        panels = [dict(label=q["letter"], graphic=q["graphic"], ident=q["ident"], landscape="landscape" in q["paper"])
                  for q in f["panels"]]
        caption = [[(f"Fig. S{k} ", True, False)] + md_parts(f["legend"][0])] + [md_parts(x) for x in f["legend"][1:]]
        add_figures(doc, f"Fig. S{k}", panels, caption, work, f"figS{k}")
    doc.save(OUT / DOCX_CITED)
    print("wrote", OUT / DOCX_CITED)


def build_extra_docx(work):
    rec = capture(BX, BX.build_pdf)
    assert len(rec["panels"]) == len(FIG_NAME)
    doc = new_doc(f"{C.TITLE_PLAIN} - raw phenotype checks (figures)")
    overview = re.sub(r" \(Figs?\. E[^)]*\)", "", rec["legend"][0])
    overview = re.sub(r" Tables: ESM_extra_phenotype_checks\.xlsx$", " The tables are in the accompanying Excel file.", overview)
    title_page(doc, "Raw phenotype checks: figures", [unnumber(overview)],
               [[(name, False, False)] for name in FIG_NAME.values()])
    for (lab, *_), pnl in zip(BX.FIGS, rec["panels"]):
        e = lab.split()[-1]
        cap = pnl["ident"]
        assert cap.startswith(f"**Fig. {e}** "), cap
        body = unnumber(cap[len(f"**Fig. {e}** "):])
        pnl = dict(pnl, ident="", label="")
        add_figures(doc, FIG_NAME[e], [pnl], [[(f"{FIG_NAME[e]}. ", True, False)] + md_parts(body)], work, f"extra_{e}")
    doc.save(OUT / DOCX_EXTRA)
    print("wrote", OUT / DOCX_EXTRA)


# ---- final checks ---------------------------------------------------------------------------------------------------
def docx_text(path):
    d = Document(str(path))
    return "\n".join(p.text for p in d.paragraphs)


def xlsx_text(path):
    wb = openpyxl.load_workbook(path)
    return "\n".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None) \
        + "\n" + "\n".join(wb.sheetnames)


def final_checks():
    for f in (XLSX_EXTRA, DOCX_EXTRA):                      # no numbers, no "not cited" statement in the unnumbered files
        t = (xlsx_text if f.endswith("xlsx") else docx_text)(OUT / f)
        assert not re.search(r"\bE\d\b|ESM_extra|not cited|unnumbered|uncited", t), f
    t = docx_text(OUT / DOCX_CITED)
    C.check_supp_refs(t, DOCX_CITED)
    for k in range(1, C.N_FIGURES + 1):
        assert f"Fig. S{k} " in t, k
    t = xlsx_text(OUT / XLSX_CITED)
    C.check_supp_refs(t, XLSX_CITED)
    assert openpyxl.load_workbook(OUT / XLSX_CITED).sheetnames == ["Contents"] + [f"Table S{k}" for k in range(1, C.N_TABLES + 1)]


if __name__ == "__main__":
    C.check_title()
    assert C.TABLES_FILE.exists(), "run build_ESM_tables.py first (build_all.sh)"
    if OUT.exists():
        shutil.rmtree(OUT)                                  # only ever holds files written by this script
    OUT.mkdir()
    work = pathlib.Path(tempfile.mkdtemp(prefix="esm_review_", dir="/mnt/data/shahar/.tmp"))
    shutil.copy2(C.TABLES_FILE, OUT / XLSX_CITED)
    print("wrote", OUT / XLSX_CITED, f"(copy of {C.TABLES_FILE.name})")
    build_extra_xlsx()
    build_cited_docx(work)
    build_extra_docx(work)
    shutil.rmtree(work)
    final_checks()
    with zipfile.ZipFile(OUT / ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for f in (XLSX_CITED, DOCX_CITED, XLSX_EXTRA, DOCX_EXTRA):
            z.write(OUT / f, f)
    print(f"zip {OUT / ZIP} ({round((OUT / ZIP).stat().st_size / 1e6, 1)} MB)")
