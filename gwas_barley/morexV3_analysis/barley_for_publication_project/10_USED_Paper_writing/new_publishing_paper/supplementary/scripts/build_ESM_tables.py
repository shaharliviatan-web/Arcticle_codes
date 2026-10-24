#!/usr/bin/env python3
"""build_ESM_tables.py -- ESM_1.xlsx, the Supplementary Tables S1-S22 of the TAG manuscript: one sheet per table,
numbered by first citation in the manuscript (S. Hübner's comment, user decision 2026-10-17; layout as the supplementary
tables of Potapenko et al. 2026a, Mol Biol Evol).

Sheet 1 "Contents": the TAG title block (article title, journal, authors, corresponding author's affiliation and
e-mail), then one row per table (number, caption, link to its sheet). Sheets "Table S1" ... "Table S22": the caption
(bold number), the description of the table, the definitions of its columns, then the table itself (header row,
autofilter), copied cell by cell from the intermediate workbooks ../parts/part_N.xlsx written by build_ESM_xlsx.py.
The map (part, sheet) -> Table S number and the captions are SUPP_TABLES in esm_common.py; references between sheets
("sheet “Loci”") are rewritten as table numbers ("Table S8").

Checks (the script stops if one fails): every table copied unchanged; no "Online Resource" left; every Table S / Fig. S
reference exists; the relevance-screen categories absent; the manuscript cites Tables S1-S22 and Figs. S1-S5, each
first cited in numerical order, and no "Online Resource" remains in its visible text.
The same workbook is the review copy for S. Hübner (package_for_review.py copies it as Supplementary_Tables.xlsx).
Run: python3 build_ESM_tables.py   (after build_ESM_xlsx.py; build_all.sh). Created 2026-10-17.
"""
import re
import openpyxl
import esm_common as C
import esm_assemble as A

FIGURES_NOTE = "The figures (Figs. S1–S5) are in the accompanying file of Supplementary Figures."


def repeats(desc, caption):
    """True when the sheet description adds nothing to the caption (>= 75% of its words are in the caption): many
    former descriptions restate what the former group caption, now split into the table captions, already says"""
    w = lambda x: set(re.findall(r"[a-zα-ωε²0-9]+", A.plain(x).lower())) - {"the", "of", "and", "a", "in", "with", "per", "each"}
    if desc.startswith("One row"):                          # the row structure is always kept
        return False
    d = w(desc)
    return bool(d) and len(d & w(caption)) / len(d) >= 0.75


def manuscript_citation_order():
    """visible text of the manuscript (hidden notes and Word-comment bodies removed): the order in which Table S and
    Fig. S numbers are first cited"""
    t = C.WORKING_FILE.read_text(encoding="utf-8")
    # from the title to the end of the text: the open-items list above the title is deleted before submission
    t = t[t.index("**Genome-wide association and haplotype analysis identify"):t.index("# Supplementary Information")]
    t = re.sub(r"<!--.*?-->", "", t, flags=re.S)
    t = re.sub(r"\[[^\[\]]*\]\{\.comment-start[^}]*\}", "", t)
    assert "Online Resource" not in t, "the manuscript still cites an Online Resource"
    order = {"Table": [], "Fig": []}
    for kind, body in re.findall(r"\b(Tables?|Figs?)\.? (S\d+[a-z]?(?:(?:, | and |–)(?:S\d+[a-z]?|[a-z]\b))*)", t):
        kind = "Table" if kind.startswith("Table") else "Fig"
        nums = []
        for a, b in re.findall(r"S(\d+)(?:[a-z]?–S(\d+))?", body):
            nums += list(range(int(a), int(b) + 1)) if b else [int(a)]
        for k in nums:
            if k not in order[kind]:
                order[kind].append(k)
    return order


def check_manuscript():
    order = manuscript_citation_order()
    assert order["Table"] == list(range(1, C.N_TABLES + 1)), f"tables first cited in the order {order['Table']}"
    assert order["Fig"] == list(range(1, C.N_FIGURES + 1)), f"figures first cited in the order {order['Fig']}"
    print(f"  ✓ manuscript: Tables S1–S{C.N_TABLES} and Figs. S1–S{C.N_FIGURES} all cited, each first cited in order; "
          f"no Online Resource left")


def build():
    C.check_title()
    wb = openpyxl.Workbook()
    parts, contents = {}, []
    for k, (part, sheet, caption) in enumerate(C.SUPP_TABLES, 1):
        if part not in parts:
            parts[part] = A.parse_about(C.PARTS / f"part_{part}.xlsx")
        src, _label, _caption, sheets, defs, notes = parts[part]
        desc = dict(sheets)[sheet]
        fix = lambda s: C.sheet_refs_to_tables(s, part)
        legend = [([(f"Table S{k} ", True, False)] + A.md_parts(caption), False)]
        if not repeats(desc, caption):                       # the sheet description, unless it only repeats the caption
            legend += [([(fix(desc), False, False)], False)]
        if defs.get(sheet):
            legend += [([("Columns", True, False)], True)]
            legend += [([(f"{col}: ", True, False), (fix(d), False, False)], False) for col, d in defs[sheet]]
        if notes and sheet == sheets[0][0]:                  # notes of a part go with its first table
            legend += [([("Notes", True, False)], True)] + [([(fix(x), False, False)], False) for x in notes]
        A.copy_table_sheet(wb, src[sheet], f"Table S{k}", legend)
        contents.append((f"Table S{k}", A.md_parts(caption), [(f"Table S{k}", None)]))
    A.contents_sheet(wb, "Supplementary Tables", [
        f"Supplementary Tables S1–S{C.N_TABLES} of the manuscript, one sheet per table, numbered in the order in which "
        f"they are first cited. Each sheet opens with the caption of the table, a description and the definitions of "
        f"its columns; the table follows. " + FIGURES_NOTE], contents, col_b="Caption")
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) or v is not None and type(v).__name__ == "CellRichText":
                    C.check_supp_refs(str(v), ws.title)
                    assert "sheet “" not in str(v), (ws.title, v)
    C.check_no_screen_categories(wb)
    wb.save(C.TABLES_FILE)
    print(f"wrote {C.TABLES_FILE}  ({len(wb.sheetnames) - 1} tables)")


if __name__ == "__main__":
    build()
    check_manuscript() if "--no-manuscript-check" not in __import__("sys").argv else None
